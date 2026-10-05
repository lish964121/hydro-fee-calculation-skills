#!/usr/bin/env python3
"""Generate a flood-related engineering design fee report as a Word document."""

from __future__ import annotations

import argparse
import json
import math
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


DEFAULT_MIN_FEE = 15000

STAGE_FACTORS = {
    "scheme": 0.35,
    "preliminary": 0.55,
    "construction_drawing": 0.75,
    "full_process": 1.00,
}

COMPLEXITY_FACTORS = {
    "low": 0.85,
    "normal": 1.00,
    "high": 1.20,
    "very_high": 1.45,
}

ENGINEERING_FACTORS = {
    "river_training": 1.10,
    "revetment": 1.05,
    "retaining_wall": 1.15,
    "bridge_culvert": 1.20,
    "pipeline": 1.15,
    "landscape_path": 1.00,
    "temporary_diversion": 1.10,
    "mixed": 1.25,
}

SPECIAL_ITEM_FEES = {
    "survey_coordination": 3000,
    "cad_drawings": 5000,
    "quantity_list": 4000,
    "cost_estimate": 5000,
    "construction_method": 3000,
    "review_response": 3000,
    "site_adjustment": 5000,
}

REQUIRED_FIELDS = [
    "project_name",
    "location",
    "design_stage",
    "engineering_type",
    "construction_cost_yuan",
]


def money(value: float) -> str:
    return f"{value:,.0f} 元"


def base_rate(construction_cost_yuan: float | None) -> float:
    if construction_cost_yuan is None:
        return 0.05
    if construction_cost_yuan <= 500000:
        return 0.060
    if construction_cost_yuan <= 2000000:
        return 0.050
    if construction_cost_yuan <= 10000000:
        return 0.040
    if construction_cost_yuan <= 50000000:
        return 0.032
    return 0.026


def get_factor(table: dict[str, float], key: str | None, default: float = 1.0) -> float:
    return table.get(str(key), default) if key is not None else default


def round_up(value: float, unit: int) -> int:
    return int(math.ceil(value / unit) * unit) if unit > 0 else int(round(value))


def calculate_fee(data: dict[str, Any]) -> dict[str, Any]:
    missing = [field for field in REQUIRED_FIELDS if data.get(field) in (None, "")]

    construction_cost = data.get("construction_cost_yuan")
    construction_cost = float(construction_cost) if construction_cost is not None else None
    rate = base_rate(construction_cost)
    stage = get_factor(STAGE_FACTORS, data.get("design_stage"), 1.0)
    complexity = get_factor(COMPLEXITY_FACTORS, data.get("complexity"), 1.0)
    engineering = get_factor(ENGINEERING_FACTORS, data.get("engineering_type"), 1.0)
    review_link = 1.10 if bool(data.get("flood_review_linked", False)) else 1.00
    discount = float(data.get("discount", 1.0))
    min_fee = float(data.get("min_fee_yuan", DEFAULT_MIN_FEE))

    base_design_fee = (construction_cost or 0) * rate
    adjusted_design_fee = base_design_fee * stage * complexity * engineering * review_link

    special_items = data.get("special_items", {})
    special_breakdown = []
    special_fee = 0.0
    for item in data.get("deliverables", []):
        item_fee = float(special_items.get(item, SPECIAL_ITEM_FEES.get(item, 0)))
        if item_fee:
            special_fee += item_fee
            special_breakdown.append((item, item_fee))

    subtotal = max(min_fee, adjusted_design_fee + special_fee)
    recommended = round_up(subtotal * discount, int(data.get("round_to_yuan", 1000)))

    return {
        "missing": missing,
        "status": "暂估测算" if missing else "正式测算",
        "construction_cost": construction_cost,
        "base_rate": rate,
        "base_design_fee": base_design_fee,
        "stage_factor": stage,
        "complexity_factor": complexity,
        "engineering_factor": engineering,
        "review_link_factor": review_link,
        "special_breakdown": special_breakdown,
        "special_fee": special_fee,
        "adjusted_design_fee": adjusted_design_fee,
        "min_fee": min_fee,
        "subtotal": subtotal,
        "discount": discount,
        "recommended_fee": recommended,
    }


def set_doc_style(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "仿宋"
    normal.font.size = Pt(12)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value


def build_docx(data: dict[str, Any], result: dict[str, Any], output: Path) -> None:
    doc = Document()
    set_doc_style(doc)
    title = doc.add_heading("防洪影响相关设计费计算说明", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"测算状态：{result['status']}")

    doc.add_heading("一、项目基本情况", level=1)
    add_table(doc, ["项目", "内容"], [
        ["项目名称", str(data.get("project_name", "待核实"))],
        ["建设单位", str(data.get("construction_unit", "待核实"))],
        ["工程地点", str(data.get("location", "待核实"))],
        ["设计阶段", str(data.get("design_stage", "待核实"))],
        ["工程类型", str(data.get("engineering_type", "待核实"))],
        ["建安投资", money(result["construction_cost"] or 0) if result["construction_cost"] else "待核实"],
        ["是否与防洪评价审查联动", "是" if data.get("flood_review_linked") else "否"],
    ])

    doc.add_heading("二、设计工作范围", level=1)
    doc.add_paragraph(
        "本次设计费适用于防洪影响评价相关的工程设计、补救措施设计、涉河工程方案或施工图设计等工作。"
        "报告编制费、模型分析费如已另列，不应与本项重复计取。"
    )

    doc.add_heading("三、计算口径", level=1)
    doc.add_paragraph(
        "设计费 = max(最低收费, 建安投资 × 基准费率 × 阶段系数 × 复杂系数 × 专业系数 × 审查联动系数 + 专项工作费) × 折扣系数。"
    )

    doc.add_heading("四、参数取值", level=1)
    add_table(doc, ["参数", "取值"], [
        ["建安投资", money(result["construction_cost"] or 0) if result["construction_cost"] else "待核实"],
        ["基准费率", f"{result['base_rate'] * 100:.2f}%"],
        ["阶段系数", f"{result['stage_factor']:.2f}"],
        ["复杂系数", f"{result['complexity_factor']:.2f}"],
        ["专业系数", f"{result['engineering_factor']:.2f}"],
        ["审查联动系数", f"{result['review_link_factor']:.2f}"],
        ["折扣系数", f"{result['discount']:.2f}"],
    ])

    doc.add_heading("五、费用计算", level=1)
    add_table(doc, ["费用项", "金额"], [
        ["基准设计费", money(result["base_design_fee"])],
        ["修正后设计费", money(result["adjusted_design_fee"])],
        ["专项工作费", money(result["special_fee"])],
        ["最低收费", money(result["min_fee"])],
        ["折扣前费用", money(result["subtotal"])],
        ["建议设计费", money(result["recommended_fee"])],
    ])

    if result["special_breakdown"]:
        doc.add_paragraph("专项工作费明细如下：")
        add_table(doc, ["专项", "金额"], [[name, money(fee)] for name, fee in result["special_breakdown"]])

    doc.add_heading("六、报价建议", level=1)
    doc.add_paragraph(
        f"本项目防洪影响相关设计费建议报价为 {money(result['recommended_fee'])}。"
        "实际合同金额应结合设计深度、审查意见回复、施工图修改和现场服务要求协商确定。"
    )

    if data.get("notes"):
        doc.add_heading("七、补充说明", level=1)
        for note in data["notes"]:
            doc.add_paragraph(str(note))

    if result["missing"]:
        doc.add_heading("八、需核实资料", level=1)
        for item in result["missing"]:
            doc.add_paragraph(str(item))

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    data = json.loads(Path(args.input).read_text(encoding="utf-8"))
    result = calculate_fee(data)
    build_docx(data, result, Path(args.output))
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
