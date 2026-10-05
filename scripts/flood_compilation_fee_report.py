#!/usr/bin/env python3
"""Generate a flood impact evaluation compilation fee report as a Word document."""

from __future__ import annotations

import argparse
import json
import math
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


BASE_FEE = 30000
DEFAULT_MIN_FEE = 25000

COMPLEXITY_FACTORS = {
    "low": 0.85,
    "normal": 1.00,
    "high": 1.25,
    "very_high": 1.50,
}

RIVER_FACTORS = {
    "一般河道": 0.90,
    "区管": 1.00,
    "市管": 1.15,
    "省管": 1.30,
    "山洪沟": 1.05,
}

MODEL_FACTORS = {
    "none": 0.85,
    "simple": 1.00,
    "one_d": 1.20,
    "two_d": 1.45,
}

REVIEW_FACTORS = {
    "none": 0.95,
    "normal": 1.00,
    "multiple_rounds": 1.15,
    "expert_meeting": 1.25,
}

SPECIAL_ITEM_FEES = {
    "drawings": 5000,
    "gis": 6000,
    "hec_ras": 8000,
    "two_d_model_files": 15000,
    "ppt": 3000,
    "meeting_minutes": 2000,
    "site_survey_extra": 5000,
}

REQUIRED_FIELDS = ["project_name", "location", "project_type", "river_level", "hydraulic_model"]


def money(value: float) -> str:
    return f"{value:,.0f} 元"


def length_factor(length_km: float | None) -> float:
    if length_km is None:
        return 1.00
    if length_km <= 0.5:
        return 0.90
    if length_km <= 2:
        return 1.00
    if length_km <= 5:
        return 1.20
    return 1.40


def area_factor(area_ha: float | None) -> float:
    if area_ha is None:
        return 1.00
    if area_ha <= 5:
        return 0.90
    if area_ha <= 20:
        return 1.00
    if area_ha <= 50:
        return 1.20
    return 1.35


def round_up(value: float, unit: int) -> int:
    return int(math.ceil(value / unit) * unit) if unit > 0 else int(round(value))


def get_factor(table: dict[str, float], key: str | None, default: float = 1.0) -> float:
    return table.get(str(key), default) if key is not None else default


def calculate_fee(data: dict[str, Any]) -> dict[str, Any]:
    missing = [field for field in REQUIRED_FIELDS if not data.get(field)]
    if data.get("river_length_km") is None and data.get("area_ha") is None:
        missing.append("river_length_km 或 area_ha")

    lf = length_factor(data.get("river_length_km"))
    af = area_factor(data.get("area_ha"))
    scale = max(lf, af)
    complexity = get_factor(COMPLEXITY_FACTORS, data.get("complexity"), 1.0)
    river = get_factor(RIVER_FACTORS, data.get("river_level"), 1.0)
    model = get_factor(MODEL_FACTORS, data.get("hydraulic_model"), 1.0)
    review = get_factor(REVIEW_FACTORS, data.get("review_level"), 1.0)
    discount = float(data.get("discount", 1.0))
    min_fee = float(data.get("min_fee_yuan", DEFAULT_MIN_FEE))

    special_items = data.get("special_items", {})
    special_breakdown = []
    special_fee = 0.0
    for item in data.get("deliverables", []):
        item_fee = float(special_items.get(item, SPECIAL_ITEM_FEES.get(item, 0)))
        if item_fee:
            special_fee += item_fee
            special_breakdown.append((item, item_fee))

    adjusted_base = BASE_FEE * scale * complexity * river * model * review
    subtotal = max(min_fee, adjusted_base + special_fee)
    recommended = round_up(subtotal * discount, int(data.get("round_to_yuan", 1000)))

    return {
        "missing": missing,
        "status": "暂估测算" if missing else "正式测算",
        "base_fee": BASE_FEE,
        "length_factor": lf,
        "area_factor": af,
        "scale_factor": scale,
        "complexity_factor": complexity,
        "river_factor": river,
        "model_factor": model,
        "review_factor": review,
        "special_breakdown": special_breakdown,
        "special_fee": special_fee,
        "adjusted_base": adjusted_base,
        "min_fee": min_fee,
        "subtotal": subtotal,
        "discount": discount,
        "recommended_fee": recommended,
    }


def set_doc_style(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "仿宋"
    normal.font.size = Pt(12)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value


def build_docx(data: dict[str, Any], result: dict[str, Any], output: Path) -> None:
    doc = Document()
    set_doc_style(doc)
    title = doc.add_heading("防洪影响评价编制费计算说明", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"测算状态：{result['status']}")

    doc.add_heading("一、项目基本情况", level=1)
    add_table(doc, ["项目", "内容"], [
        ["项目名称", str(data.get("project_name", "待核实"))],
        ["建设单位", str(data.get("construction_unit", "待核实"))],
        ["工程地点", str(data.get("location", "待核实"))],
        ["工程类型", str(data.get("project_type", "待核实"))],
        ["河道等级", str(data.get("river_level", "待核实"))],
        ["防洪标准", str(data.get("design_flood_standard", "待核实"))],
        ["涉河长度", f"{data.get('river_length_km', '待核实')} km"],
        ["评价范围面积", f"{data.get('area_ha', '待核实')} hm2"],
        ["水力模型深度", str(data.get("hydraulic_model", "待核实"))],
    ])

    doc.add_heading("二、计算口径", level=1)
    doc.add_paragraph(
        "本次费用测算采用“基础编制费 + 专项工作费”的工程咨询报价口径。"
        "如业主、地方主管部门或公司已有固定报价标准，应优先采用固定标准。"
    )
    doc.add_paragraph(
        "编制费 = max(最低收费, 基础编制费 × 规模系数 × 复杂系数 × 河道系数 × 模型系数 × 评审系数 + 专项工作费) × 折扣系数。"
    )

    doc.add_heading("三、参数取值", level=1)
    add_table(doc, ["参数", "取值"], [
        ["基础编制费", money(result["base_fee"])],
        ["长度系数", f"{result['length_factor']:.2f}"],
        ["面积系数", f"{result['area_factor']:.2f}"],
        ["规模系数", f"{result['scale_factor']:.2f}"],
        ["复杂系数", f"{result['complexity_factor']:.2f}"],
        ["河道系数", f"{result['river_factor']:.2f}"],
        ["模型系数", f"{result['model_factor']:.2f}"],
        ["评审系数", f"{result['review_factor']:.2f}"],
        ["折扣系数", f"{result['discount']:.2f}"],
    ])

    doc.add_heading("四、费用计算", level=1)
    add_table(doc, ["费用项", "金额"], [
        ["修正后基础费用", money(result["adjusted_base"])],
        ["专项工作费", money(result["special_fee"])],
        ["最低收费", money(result["min_fee"])],
        ["折扣前费用", money(result["subtotal"])],
        ["建议编制费", money(result["recommended_fee"])],
    ])

    if result["special_breakdown"]:
        doc.add_paragraph("专项工作费明细如下：")
        add_table(doc, ["专项", "金额"], [[name, money(fee)] for name, fee in result["special_breakdown"]])

    doc.add_heading("五、报价建议", level=1)
    doc.add_paragraph(
        f"本项目防洪影响评价报告编制费建议报价为 {money(result['recommended_fee'])}。"
        "实际合同金额可结合评审深度、模型成果交付要求和后续补充工作量协商确定。"
    )

    if data.get("notes"):
        doc.add_heading("六、补充说明", level=1)
        for note in data["notes"]:
            doc.add_paragraph(str(note))

    if result["missing"]:
        doc.add_heading("七、需核实资料", level=1)
        for item in result["missing"]:
            doc.add_paragraph(str(item))

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    data = json.loads(Path(args.input).read_text(encoding="utf-8"))
    result = calculate_fee(data)
    build_docx(data, result, Path(args.output))
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
